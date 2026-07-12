"""Tests for built-in tool safety and schemas."""

import asyncio
import sys
import zipfile
from pathlib import Path
from types import SimpleNamespace

import pytest

import magent.tools as tools_module
import magent.tools.executor as executor_module
from magent.permissions import RiskTier, classify_file_op, classify_shell_command
from magent.tools import ToolExecutor


def test_file_read_outside_cwd_requires_confirmation(tmp_path: Path) -> None:
    outside = tmp_path.parent / "outside.txt"
    assert classify_file_op("read", str(outside), str(tmp_path)) == RiskTier.CONFIRM


@pytest.mark.asyncio
async def test_outline_file_reports_python_symbols(tmp_path: Path) -> None:
    source = tmp_path / "demo.py"
    source.write_text(
        "class Demo:\n"
        "    def method(self):\n"
        "        pass\n\n"
        "async def run():\n"
        "    pass\n",
        encoding="utf-8",
    )
    tools = ToolExecutor(str(tmp_path))

    result = await tools.outline_file("demo.py")

    assert result["ok"] is True
    names = {symbol["name"] for symbol in result["symbols"]}
    assert {"Demo", "method", "run"} <= names


def test_shell_control_read_only_chains_do_not_prompt_spam() -> None:
    assert classify_shell_command("cat file.html 2>&1 | wc -l") == RiskTier.SILENT
    assert classify_shell_command("wc -l file.html 2>/dev/null; echo ---; tail -5 file.html") == RiskTier.SILENT
    assert classify_shell_command("pip list 2>/dev/null | grep -iE 'docx|pptx|python'") == RiskTier.SILENT
    assert classify_shell_command('cd /tmp && python3 -c "import docx; print(docx.__version__)"') == RiskTier.SILENT
    assert classify_shell_command("curl -s 'https://example.com' | head -500") == RiskTier.AUTO
    assert (
        classify_shell_command(
            "curl -s 'https://example.com' | grep -i 'primary\\|color\\|font' | head -100"
        )
        == RiskTier.AUTO
    )
    assert classify_shell_command("python3 -c 'print(\"hello\")'") == RiskTier.SILENT


def test_network_fetch_mutation_or_download_still_requires_confirmation() -> None:
    assert classify_shell_command("curl -X POST https://example.com") == RiskTier.CONFIRM
    assert classify_shell_command("curl -s https://example.com -o page.html") == RiskTier.CONFIRM
    assert classify_shell_command("wget https://example.com -O page.html") == RiskTier.CONFIRM


def test_macos_shell_rewrites_prefer_python3(monkeypatch) -> None:
    monkeypatch.setattr(executor_module, "sys", SimpleNamespace(platform="darwin"))

    assert (
        executor_module._prefer_platform_python_command("pip install python-pptx")
        == "python3 -m pip install python-pptx"
    )
    assert (
        executor_module._prefer_platform_python_command("pip list | grep docx")
        == "python3 -m pip list | grep docx"
    )
    assert executor_module._prefer_platform_python_command("python -c 'print(1)'") == "python3 -c 'print(1)'"


def test_shell_control_blocks_dangerous_segment_even_if_allowlisted() -> None:
    assert classify_shell_command("git status; rm -rf build", ["git *"]) == RiskTier.BLOCK


@pytest.mark.asyncio
async def test_noninteractive_permissions_return_structured_denial(tmp_path: Path) -> None:
    tools = ToolExecutor(
        str(tmp_path),
        permission_mode="balanced",
        interactive_permissions=False,
    )

    result = await tools.run_shell("git status; rm -rf build")

    assert result["ok"] is False
    assert result["permission_required"] is True
    assert result["permission_reason"] == "permission-required"


@pytest.mark.asyncio
async def test_web_search_prefers_relevant_ddgs_results(monkeypatch, tmp_path: Path) -> None:
    class FakeDDGS:
        def __enter__(self):
            return self

        def __exit__(self, *args):
            return None

        def text(self, query: str, max_results: int = 8):
            return [
                {
                    "title": "complete - German translation",
                    "body": "Dictionary entry for complete.",
                    "href": "https://dict.example/complete",
                },
                {
                    "title": "A History of Cheese",
                    "body": "Cheese making has ancient origins and a long timeline.",
                    "href": "https://example.com/history-of-cheese",
                },
                {
                    "title": "The History of Cheese: From Ancient Origins to Modern Day",
                    "body": "Short social video caption.",
                    "href": "https://www.tiktok.com/@example/video/123",
                },
            ]

    monkeypatch.setitem(sys.modules, "ddgs", SimpleNamespace(DDGS=FakeDDGS))
    tools = ToolExecutor(str(tmp_path), permission_mode="silent")

    result = await tools.web_search("Complete history of cheese from ancient origins", max_results=8)

    assert result["ok"] is True
    assert result["source"] == "ddgs"
    assert result["filtered_count"] == 2
    assert [item["url"] for item in result["results"]] == ["https://example.com/history-of-cheese"]


def test_tool_definitions_have_required_arguments(tmp_path: Path) -> None:
    tools = ToolExecutor(str(tmp_path))
    defs = {d["function"]["name"]: d for d in tools.get_tool_definitions()}

    read_file_required = defs["read_file"]["function"]["parameters"]["required"]
    write_file_required = defs["write_file"]["function"]["parameters"]["required"]
    write_file_props = defs["write_file"]["function"]["parameters"]["properties"]

    assert "path" in read_file_required
    assert {"path", "content"} <= set(write_file_required)
    assert "activity" in write_file_props
    assert "activity" not in write_file_required
    assert {"phase", "intent", "expected"} <= set(write_file_props["activity"]["properties"])
    assert "outline_file" in defs
    assert "read_file_range" in defs
    assert "deep_research" in defs
    assert "create_docx" in defs
    assert "create_pptx" in defs
    assert "create_svg" in defs
    assert "create_diagram" in defs
    assert "create_image" in defs
    assert "generate_image" in defs


@pytest.mark.asyncio
async def test_dispatch_reports_missing_required_arguments(tmp_path: Path) -> None:
    tools = ToolExecutor(str(tmp_path), permission_mode="silent")

    result = await tools.dispatch("write_file", {"path": "missing-content.txt"})

    assert result["ok"] is False
    assert result["missing"] == ["content"]


@pytest.mark.asyncio
async def test_write_file_accepts_common_content_aliases(tmp_path: Path) -> None:
    tools = ToolExecutor(str(tmp_path), permission_mode="yolo")

    html = await tools.dispatch("write_file", {"path": "page.html", "html": "<!doctype html><h1>ok</h1>"})
    file_content = await tools.dispatch(
        "write_file",
        {
            "filename": "notes.md",
            "file_content": "# Notes\n\nComplete content.",
        },
    )

    assert html["ok"] is True
    assert file_content["ok"] is True
    assert (tmp_path / "page.html").read_text(encoding="utf-8") == "<!doctype html><h1>ok</h1>"
    assert (tmp_path / "notes.md").read_text(encoding="utf-8") == "# Notes\n\nComplete content."


@pytest.mark.asyncio
async def test_document_tools_create_docx_and_pptx(tmp_path: Path) -> None:
    tools = ToolExecutor(str(tmp_path), permission_mode="silent")
    sections = [
        {
            "title": "Origins",
            "content": "Oranges emerged from ancient citrus cultivation.",
            "bullets": ["Hybrid of pomelo and mandarin", "Spread through trade routes"],
        },
        {"title": "Modern Impact", "bullets": ["Global crop", "Vitamin C source"]},
    ]

    docx = await tools.create_docx(
        "history_of_oranges.docx",
        "History of Oranges",
        sections,
        "A concise research brief",
    )
    pptx = await tools.create_pptx(
        "history_of_oranges.pptx",
        "History of Oranges",
        sections,
        "From citrus groves to global tables",
    )

    assert docx["ok"] is True
    assert pptx["ok"] is True
    assert (tmp_path / "history_of_oranges.docx").stat().st_size > 0
    assert (tmp_path / "history_of_oranges.pptx").stat().st_size > 0

    from docx import Document
    from pptx import Presentation

    document_text = "\n".join(paragraph.text for paragraph in Document(tmp_path / "history_of_oranges.docx").paragraphs)
    presentation = Presentation(tmp_path / "history_of_oranges.pptx")

    assert "History of Oranges" in document_text
    assert "Origins" in document_text
    assert len(presentation.slides) == 3


@pytest.mark.asyncio
async def test_visual_artifact_tools_create_svg_diagram_and_image(tmp_path: Path) -> None:
    tools = ToolExecutor(str(tmp_path), permission_mode="silent")
    elements = [
        {"type": "rect", "x": 20, "y": 20, "width": 180, "height": 90, "fill": "#f97316"},
        {"type": "text", "x": 42, "y": 72, "text": "Orange", "fill": "#ffffff"},
    ]

    svg = await tools.create_svg("orange.svg", elements, "Orange card", 240, 140)
    diagram = await tools.create_diagram(
        "orange-flow.mmd",
        "Orange Journey",
        [{"id": "origin", "label": "Origins"}, {"id": "trade", "label": "Trade"}],
        [{"from": "origin", "to": "trade", "label": "spreads"}],
        "LR",
    )
    image = await tools.create_image("orange.png", elements, "Orange card", 240, 140, "#fff7ed")

    assert svg["ok"] is True
    assert diagram["ok"] is True
    assert image["ok"] is True
    assert "<svg" in (tmp_path / "orange.svg").read_text(encoding="utf-8")
    assert "flowchart LR" in (tmp_path / "orange-flow.mmd").read_text(encoding="utf-8")

    from PIL import Image

    rendered = Image.open(tmp_path / "orange.png")
    assert rendered.size == (240, 140)


@pytest.mark.asyncio
async def test_generate_image_uses_image_maker_role(monkeypatch, tmp_path: Path) -> None:
    from magent.config import Config

    async def fake_aimage_generation(**kwargs):
        assert kwargs["model"] == "gpt-image-1"
        assert kwargs["prompt"] == "flat vector navy blue iceberg"
        return SimpleNamespace(
            data=[SimpleNamespace(b64_json="iVBORw0KGgo=")],
        )

    monkeypatch.setitem(
        sys.modules,
        "litellm",
        SimpleNamespace(aimage_generation=fake_aimage_generation, suppress_debug_info=False),
    )
    cfg = Config(
        {
            "defaults": {"provider": "ollama", "model": "qwen"},
            "models": {"image_maker": "openai/gpt-image-1"},
            "providers": {"openai": {"api_key": "key"}},
        }
    )
    tools = ToolExecutor(str(tmp_path), permission_mode="silent", config=cfg)

    result = await tools.generate_image("iceberg.png", "flat vector navy blue iceberg")

    assert result["ok"] is True
    assert (tmp_path / "iceberg.png").read_bytes() == b"\x89PNG\r\n\x1a\n"


@pytest.mark.asyncio
async def test_file_tools_read_write_edit_list_diff_and_range(tmp_path: Path) -> None:
    tools = ToolExecutor(str(tmp_path), permission_mode="silent")

    written = await tools.write_file("notes/a.txt", "one\ntwo\nthree\n")
    read = await tools.read_file("notes/a.txt")
    ranged = await tools.read_file_range("notes/a.txt", 2, 3)
    edited = await tools.edit_file("notes/a.txt", "two", "TWO")
    (tmp_path / "notes" / "b.txt").write_text("one\nTWO\nfour\n", encoding="utf-8")
    diff = await tools.diff_files("notes/a.txt", "notes/b.txt")
    listed = await tools.list_dir("notes")

    assert written["ok"] is True
    assert read["lines"] == 3
    assert "2: two" in ranged["content"]
    assert edited["ok"] is True
    assert diff["changed"] is True
    assert {entry["name"] for entry in listed["entries"]} == {"a.txt", "b.txt"}


@pytest.mark.asyncio
async def test_write_file_rejects_filename_as_content(tmp_path: Path) -> None:
    tools = ToolExecutor(str(tmp_path), permission_mode="silent")

    result = await tools.write_file("cheese.html", "cheese.html")

    assert result["ok"] is False
    assert result["blocked_by"] == "write-file-content-guard"
    assert not (tmp_path / "cheese.html").exists()


@pytest.mark.asyncio
async def test_write_file_accepts_real_html_content(tmp_path: Path) -> None:
    tools = ToolExecutor(str(tmp_path), permission_mode="silent")
    html = "<!doctype html><html><body><main><h1>History of Cheese</h1></main></body></html>"

    result = await tools.write_file("cheese.html", html)

    assert result["ok"] is True
    assert (tmp_path / "cheese.html").read_text(encoding="utf-8") == html


@pytest.mark.asyncio
async def test_write_file_skips_identical_duplicate_content(tmp_path: Path) -> None:
    tools = ToolExecutor(str(tmp_path), permission_mode="silent")
    content = "same content\n"

    first = await tools.write_file("notes.txt", content)
    second = await tools.write_file("notes.txt", content)

    assert first["ok"] is True
    assert second["ok"] is True
    assert second["unchanged"] is True
    assert "checkpoint_id" not in second


@pytest.mark.asyncio
async def test_dispatch_normalizes_common_write_file_aliases(tmp_path: Path) -> None:
    tools = ToolExecutor(str(tmp_path), permission_mode="silent")

    result = await tools.dispatch(
        "write_file",
        {"file_path": "notes/alias.txt", "contents": "alias worked"},
    )

    assert result["ok"] is True
    assert (tmp_path / "notes" / "alias.txt").read_text(encoding="utf-8") == "alias worked"


@pytest.mark.asyncio
async def test_dispatch_strips_tool_activity_metadata(tmp_path: Path) -> None:
    tools = ToolExecutor(str(tmp_path), permission_mode="silent")

    result = await tools.dispatch(
        "write_file",
        {
            "path": "notes/activity.txt",
            "content": "metadata ignored by handler",
            "activity": {
                "phase": "edit",
                "intent": "Create the requested notes file.",
                "expected": "A file exists on disk.",
            },
        },
    )

    assert result["ok"] is True
    assert (tmp_path / "notes" / "activity.txt").read_text(encoding="utf-8") == "metadata ignored by handler"


@pytest.mark.asyncio
async def test_shell_control_can_be_session_allowed(tmp_path: Path) -> None:
    command = "echo 123 && echo 456"
    tools = ToolExecutor(
        str(tmp_path),
        permission_mode="balanced",
        trusted_shell_patterns=[command],
    )

    result = tools._check_shell_permission(command, RiskTier.CONFIRM)

    assert result.approved is True
    assert result.reason == "trusted-shell"


@pytest.mark.asyncio
async def test_run_shell_declines_native_file_writes_without_prompt(tmp_path: Path) -> None:
    tools = ToolExecutor(str(tmp_path), permission_mode="balanced", interactive_permissions=False)

    heredoc = await tools.run_shell("cat > cheese.html << 'EOF'\n<h1>Cheese</h1>\nEOF")
    python_write = await tools.run_shell(
        "python3 -c \"open('cheese.html', 'w').write('<h1>Cheese</h1>')\""
    )

    assert heredoc["ok"] is False
    assert heredoc["blocked_by"] == "native-file-tool-policy"
    assert heredoc["recommended_tool"] == "write_file"
    assert python_write["ok"] is False
    assert python_write["blocked_by"] == "native-file-tool-policy"
    assert not (tmp_path / "cheese.html").exists()


@pytest.mark.asyncio
async def test_deep_research_collects_sources(monkeypatch, tmp_path: Path) -> None:
    tools = ToolExecutor(str(tmp_path), permission_mode="silent")

    async def fake_search(query: str, max_results: int = 8):
        return {
            "ok": True,
            "query": query,
            "results": [
                {
                    "title": "Example research",
                    "snippet": "Useful source snippet",
                    "url": "https://example.com/research",
                }
            ],
        }

    async def fake_fetch(url: str, extract_article: bool = True):
        return {
            "ok": True,
            "url": url,
            "status": 200,
            "content": "Long article text about the research topic.",
            "extractor": "test",
        }

    monkeypatch.setattr(tools, "web_search", fake_search)
    monkeypatch.setattr(tools, "web_fetch", fake_fetch)

    result = await tools.deep_research("agent UX", questions=["desktop apps"], max_sources=2)

    assert result["ok"] is True
    assert result["source_count"] == 1
    assert result["sources"][0]["url"] == "https://example.com/research"
    assert "Research summary" in result["summary"]


@pytest.mark.asyncio
async def test_archive_tools_roundtrip_zip(tmp_path: Path) -> None:
    (tmp_path / "src").mkdir()
    (tmp_path / "src" / "demo.txt").write_text("hello", encoding="utf-8")
    tools = ToolExecutor(str(tmp_path), permission_mode="silent")

    compressed = await tools.compress("src", "archive.zip", format="zip")
    extracted = await tools.extract("archive.zip", "out")

    assert compressed["ok"] is True
    assert zipfile.is_zipfile(tmp_path / "archive.zip")
    assert extracted["ok"] is True
    assert (tmp_path / "out" / "src" / "demo.txt").read_text(encoding="utf-8") == "hello"


@pytest.mark.asyncio
async def test_run_python_and_search_codebase(tmp_path: Path, monkeypatch) -> None:
    (tmp_path / "app.py").write_text("needle = 'found'\n", encoding="utf-8")
    tools = ToolExecutor(str(tmp_path), permission_mode="silent")

    class FakeProc:
        returncode = 0

        async def communicate(self):
            return (b"hello from python\n", b"")

        def kill(self):
            return None

    async def fake_create_exec_process(*argv, cwd):
        return FakeProc()

    monkeypatch.setattr(executor_module, "_create_exec_process", fake_create_exec_process)

    python = await tools.run_python("print('hello from python')")
    search = await tools.search_codebase("needle")

    assert python["ok"] is True
    assert python["stdout"].strip() == "hello from python"
    assert search["ok"] is True
    assert search["total"] >= 1


@pytest.mark.asyncio
async def test_run_shell_expands_bash_brace_directories(tmp_path: Path, monkeypatch) -> None:
    command = "mkdir -p teal-blog/src/{pages/blog,layouts,styles,content/blog}"
    tools = ToolExecutor(str(tmp_path), permission_mode="silent", trusted_shell_patterns=[command])

    class FakeProc:
        returncode = 0

        async def communicate(self):
            return (b"", b"")

        def kill(self):
            return None

    async def fake_create_shell_process(shell_command, cwd):
        assert shell_command == command
        (Path(cwd) / "teal-blog" / "src" / "pages" / "blog").mkdir(parents=True)
        (Path(cwd) / "teal-blog" / "src" / "layouts").mkdir(parents=True)
        (Path(cwd) / "teal-blog" / "src" / "styles").mkdir(parents=True)
        (Path(cwd) / "teal-blog" / "src" / "content" / "blog").mkdir(parents=True)
        return FakeProc()

    monkeypatch.setattr(executor_module, "_create_shell_process", fake_create_shell_process)

    result = await tools.run_shell(command)

    assert result["ok"] is True
    assert (tmp_path / "teal-blog" / "src" / "pages" / "blog").is_dir()
    assert (tmp_path / "teal-blog" / "src" / "layouts").is_dir()
    assert (tmp_path / "teal-blog" / "src" / "styles").is_dir()
    assert (tmp_path / "teal-blog" / "src" / "content" / "blog").is_dir()
    assert not (tmp_path / "teal-blog" / "src" / "{pages").exists()


@pytest.mark.asyncio
async def test_run_shell_uses_longer_default_timeout_for_js_installs(tmp_path: Path, monkeypatch) -> None:
    tools = ToolExecutor(str(tmp_path), permission_mode="silent", trusted_shell_patterns=["npm install"])
    captured: dict[str, int] = {}

    class FakeProc:
        returncode = 0

        async def communicate(self):
            return (b"", b"")

        def kill(self):
            return None

    async def fake_wait_for(coro, timeout):
        captured["timeout"] = timeout
        return await coro

    async def fake_create_subprocess_exec(*argv, **kwargs):
        return FakeProc()

    monkeypatch.setattr(tools_module.asyncio, "create_subprocess_exec", fake_create_subprocess_exec)
    monkeypatch.setattr(tools_module.asyncio, "wait_for", fake_wait_for)

    result = await tools.run_shell("npm install")

    assert result["ok"] is True
    assert captured["timeout"] == 300


@pytest.mark.asyncio
async def test_cancel_active_stops_running_shell_task(tmp_path: Path) -> None:
    command = "sleep 30"
    tools = ToolExecutor(str(tmp_path), permission_mode="silent", trusted_shell_patterns=[command])

    task = asyncio.create_task(tools.dispatch("run_shell", {"command": command, "timeout": 60}))
    await asyncio.sleep(0.1)
    assert tools.has_active_work() is True

    await tools.cancel_active()

    with pytest.raises(asyncio.CancelledError):
        await task
    assert tools.has_active_work() is False


@pytest.mark.asyncio
async def test_json_query_and_docs_search(tmp_path: Path) -> None:
    (tmp_path / "data.json").write_text('{"items": [{"name": "Ada"}]}', encoding="utf-8")
    tools = ToolExecutor(str(tmp_path), permission_mode="silent")

    query = await tools.json_query("data.json", "items[0].name")
    docs = await tools.magent_docs_search("semantic memory", limit=2)

    assert query == {"ok": True, "result": "Ada"}
    assert docs["ok"] is True
    assert docs["results"]
